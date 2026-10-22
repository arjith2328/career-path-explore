import io
import re
import zipfile
import xml.etree.ElementTree as ET
from backend.parsers.base import BaseParser

class DOCXParser(BaseParser):
    def parse(self, file_bytes: bytes, filename: str = "document.docx") -> str:
        full_text = ""
        
        # 1. Primary: python-docx
        try:
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            paragraphs.append(cell.text.strip())
            full_text = "\n".join(paragraphs).strip()
        except ImportError:
            pass
        except Exception:
            pass

        # 2. Native XML fallback
        if not full_text:
            try:
                with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
                    xml_content = z.read('word/document.xml')
                    tree = ET.fromstring(xml_content)
                    text_elements = []
                    for node in tree.iter():
                        if node.tag.endswith('}t'):
                            if node.text:
                                text_elements.append(node.text)
                        elif node.tag.endswith('}p'):
                            text_elements.append("\n")
                    full_text = "".join(text_elements).strip()
            except Exception as e:
                raise ValueError(f"Failed to parse DOCX document '{filename}': {str(e)}")

        cleaned_text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]', '', full_text)
        cleaned_text = re.sub(r'\n{3,}', '\n\n', cleaned_text).strip()

        if not cleaned_text or len(cleaned_text.strip()) < 15:
            raise ValueError(
                "NO_TEXT_EXTRACTED: No readable text was found in this DOCX document."
            )

        return cleaned_text
